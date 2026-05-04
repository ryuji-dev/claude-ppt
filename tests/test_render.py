import json
from pathlib import Path

import pytest
from pptx import Presentation

from claude_ppt.render import render

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def sample_outline() -> dict:
    return json.loads((FIXTURES / "sample-outline.json").read_text(encoding="utf-8"))


def test_render_creates_pptx_file(tmp_path: Path, sample_outline: dict) -> None:
    output = tmp_path / "out.pptx"
    render(sample_outline, output)
    assert output.exists()
    assert output.stat().st_size > 0


def test_render_slide_count_matches_outline(tmp_path: Path, sample_outline: dict) -> None:
    output = tmp_path / "out.pptx"
    render(sample_outline, output)
    pres = Presentation(str(output))
    assert len(pres.slides) == len(sample_outline["slides"])


def test_render_slide_titles_match_outline(tmp_path: Path, sample_outline: dict) -> None:
    output = tmp_path / "out.pptx"
    render(sample_outline, output)
    pres = Presentation(str(output))
    actual_titles = [_first_text(slide) for slide in pres.slides]
    expected_titles = [s["title"] for s in sample_outline["slides"]]
    assert actual_titles == expected_titles


def test_render_unknown_layout_raises(tmp_path: Path, sample_outline: dict) -> None:
    sample_outline["slides"][0]["layout"] = "nonexistent-layout"
    output = tmp_path / "out.pptx"
    with pytest.raises(ValueError, match="nonexistent-layout"):
        render(sample_outline, output)


def _first_text(slide) -> str:
    """Return the first non-empty text frame text on the slide."""
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                return text
    return ""
