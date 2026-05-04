"""hero-cards 레이아웃 — 2~3개의 강조 카드를 가로로 배치.

content 스키마:
    {"cards": [{"emoji": "🚀", "label": "라벨", "text": "설명"}, ...]}
"""

from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches

from .. import theme
from . import _common

CARDS_TOP = Inches(2.4)
CARDS_HEIGHT = Inches(4.0)
CARDS_GAP = Inches(0.4)
CARDS_MARGIN_X = Inches(0.8)


def render(slide: Slide, content: dict, title: str) -> None:
    _common.add_title(slide, title)
    cards = content.get("cards", [])
    if not cards:
        return
    n = len(cards)
    available = (
        int(_common.SLIDE_WIDTH) - 2 * int(CARDS_MARGIN_X) - int(CARDS_GAP) * (n - 1)
    )
    card_w = available // n
    for i, card in enumerate(cards):
        left = int(CARDS_MARGIN_X) + (card_w + int(CARDS_GAP)) * i
        _draw_card(slide, left, card_w, card)


def _draw_card(slide: Slide, left: int, width: int, card: dict) -> None:
    shape = _common.add_card(
        slide,
        left=left,
        top=int(CARDS_TOP),
        width=width,
        height=int(CARDS_HEIGHT),
    )
    tf = shape.text_frame
    _common.add_paragraph(
        tf,
        card.get("emoji", ""),
        size_pt=theme.EMOJI_SIZE_PT,
        color=theme.TEXT_PRIMARY,
        align=PP_ALIGN.CENTER,
        new=False,
    )
    _common.add_paragraph(
        tf,
        card.get("label", ""),
        size_pt=theme.LABEL_SIZE_PT,
        color=theme.TEXT_SECONDARY,
        align=PP_ALIGN.CENTER,
    )
    _common.add_paragraph(
        tf,
        card.get("text", ""),
        size_pt=theme.TEXT_SIZE_PT,
        color=theme.TEXT_PRIMARY,
        align=PP_ALIGN.CENTER,
    )
