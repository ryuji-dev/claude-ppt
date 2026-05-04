"""diagram-box 레이아웃 — 중앙 노드 1개 + 하위 3~6개 노드.

content 스키마:
    {
      "center": {"label": "중앙"},
      "nodes":  [{"label": "하위1"}, {"label": "하위2"}, ...]
    }
"""

from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt

from .. import theme
from . import _common

CENTER_WIDTH = Inches(3.4)
CENTER_HEIGHT = Inches(1.4)
CENTER_TOP = Inches(2.2)

NODE_TOP = Inches(5.0)
NODE_HEIGHT = Inches(1.6)
NODE_GAP = Inches(0.3)
NODE_MARGIN_X = Inches(0.8)


def render(slide: Slide, content: dict, title: str) -> None:
    _common.add_title(slide, title)
    center = content.get("center", {})
    nodes = content.get("nodes", [])

    center_left = (int(_common.SLIDE_WIDTH) - int(CENTER_WIDTH)) // 2
    center_shape = _common.add_card(
        slide,
        left=center_left,
        top=int(CENTER_TOP),
        width=int(CENTER_WIDTH),
        height=int(CENTER_HEIGHT),
        fill=theme.ACCENT_PURPLE,
        border=theme.ACCENT_PURPLE,
    )
    _common.add_paragraph(
        center_shape.text_frame,
        center.get("label", ""),
        size_pt=theme.TEXT_SIZE_PT + 6,
        color=theme.TEXT_PRIMARY,
        bold=True,
        align=PP_ALIGN.CENTER,
        new=False,
    )

    if not nodes:
        return
    n = len(nodes)
    available = (
        int(_common.SLIDE_WIDTH)
        - 2 * int(NODE_MARGIN_X)
        - int(NODE_GAP) * (n - 1)
    )
    node_w = available // n
    center_x = center_left + int(CENTER_WIDTH) // 2
    center_y_bottom = int(CENTER_TOP) + int(CENTER_HEIGHT)
    for i, node in enumerate(nodes):
        left = int(NODE_MARGIN_X) + (node_w + int(NODE_GAP)) * i
        node_shape = _common.add_card(
            slide,
            left=left,
            top=int(NODE_TOP),
            width=node_w,
            height=int(NODE_HEIGHT),
        )
        _common.add_paragraph(
            node_shape.text_frame,
            node.get("label", ""),
            size_pt=theme.TEXT_SIZE_PT,
            color=theme.TEXT_PRIMARY,
            align=PP_ALIGN.CENTER,
            new=False,
        )
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            begin_x=center_x,
            begin_y=center_y_bottom,
            end_x=left + node_w // 2,
            end_y=int(NODE_TOP),
        )
        connector.line.color.rgb = theme.ACCENT_PURPLE
        connector.line.width = Pt(1.5)
