"""step-flow 레이아웃 — 3~5개 단계를 가로로 나열, 각 카드에 번호·제목·설명.

content 스키마:
    {"steps": [{"title": "설치", "description": "..."}, ...]}
"""

from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches

from .. import theme
from . import _common

STEPS_TOP = Inches(2.4)
STEPS_HEIGHT = Inches(4.4)
STEPS_GAP = Inches(0.3)
STEPS_MARGIN_X = Inches(0.8)


def render(slide: Slide, content: dict, title: str) -> None:
    _common.add_title(slide, title)
    steps = content.get("steps", [])
    if not steps:
        return
    n = len(steps)
    available = (
        int(_common.SLIDE_WIDTH)
        - 2 * int(STEPS_MARGIN_X)
        - int(STEPS_GAP) * (n - 1)
    )
    step_w = available // n
    for i, step in enumerate(steps):
        left = int(STEPS_MARGIN_X) + (step_w + int(STEPS_GAP)) * i
        shape = _common.add_card(
            slide,
            left=left,
            top=int(STEPS_TOP),
            width=step_w,
            height=int(STEPS_HEIGHT),
        )
        tf = shape.text_frame
        _common.add_paragraph(
            tf,
            f"STEP {i + 1:02d}",
            size_pt=theme.LABEL_SIZE_PT,
            color=theme.ACCENT_PURPLE,
            bold=True,
            align=PP_ALIGN.CENTER,
            new=False,
        )
        _common.add_paragraph(
            tf,
            step.get("title", ""),
            size_pt=theme.TEXT_SIZE_PT + 4,
            color=theme.TEXT_PRIMARY,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        _common.add_paragraph(
            tf,
            step.get("description", ""),
            size_pt=theme.TEXT_SIZE_PT - 2,
            color=theme.TEXT_SECONDARY,
            align=PP_ALIGN.CENTER,
        )
