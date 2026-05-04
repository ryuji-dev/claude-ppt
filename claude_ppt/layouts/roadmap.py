"""roadmap 레이아웃 — 3~6개 단계를 가로 진행 표로 표시.

content 스키마:
    {"steps": [{"label": "1단계"}, {"label": "2단계"}, ...]}
"""

from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches

from .. import theme
from . import _common

STEPS_TOP = Inches(3.0)
STEPS_HEIGHT = Inches(1.6)
STEPS_GAP = Inches(0.25)
STEPS_MARGIN_X = Inches(0.8)


def render(slide: Slide, content: dict, title: str) -> None:
    _common.add_title(slide, title)
    steps = content.get("steps", [])
    if not steps:
        return
    n = len(steps)
    available = (
        int(_common.SLIDE_WIDTH)
        - 2 * int(STEPS_MARGIN_X)
        - int(STEPS_GAP) * (n - 1)
    )
    step_w = available // n
    for i, step in enumerate(steps):
        left = int(STEPS_MARGIN_X) + (step_w + int(STEPS_GAP)) * i
        shape = _common.add_card(
            slide,
            left=left,
            top=int(STEPS_TOP),
            width=step_w,
            height=int(STEPS_HEIGHT),
        )
        tf = shape.text_frame
        _common.add_paragraph(
            tf,
            f"{i + 1:02d}",
            size_pt=theme.LABEL_SIZE_PT,
            color=theme.ACCENT_BLUE,
            bold=True,
            align=PP_ALIGN.CENTER,
            new=False,
        )
        _common.add_paragraph(
            tf,
            step.get("label", ""),
            size_pt=theme.TEXT_SIZE_PT,
            color=theme.TEXT_PRIMARY,
            align=PP_ALIGN.CENTER,
        )
