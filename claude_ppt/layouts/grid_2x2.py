"""grid-2x2 레이아웃 — 4~6개 카드를 격자로 배치.

content 스키마:
    {"cards": [{"emoji": "🚀", "title": "기능1", "description": "..."}, ...]}

4개 → 2x2, 5~6개 → 3x2 자동 배치.
"""

from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches

from .. import theme
from . import _common

GRID_TOP = Inches(2.2)
GRID_HEIGHT = Inches(4.8)
GRID_GAP = Inches(0.3)
GRID_MARGIN_X = Inches(0.8)


def render(slide: Slide, content: dict, title: str) -> None:
    _common.add_title(slide, title)
    cards = content.get("cards", [])
    if not cards:
        return
    n = len(cards)
    cols = 2 if n <= 4 else 3
    rows = (n + cols - 1) // cols

    available_w = (
        int(_common.SLIDE_WIDTH)
        - 2 * int(GRID_MARGIN_X)
        - int(GRID_GAP) * (cols - 1)
    )
    card_w = available_w // cols

    available_h = int(GRID_HEIGHT) - int(GRID_GAP) * (rows - 1)
    card_h = available_h // rows

    for i, card in enumerate(cards):
        row = i // cols
        col = i % cols
        left = int(GRID_MARGIN_X) + (card_w + int(GRID_GAP)) * col
        top = int(GRID_TOP) + (card_h + int(GRID_GAP)) * row
        _draw_card(slide, left, top, card_w, card_h, card)


def _draw_card(slide: Slide, left: int, top: int, w: int, h: int, card: dict) -> None:
    shape = _common.add_card(slide, left=left, top=top, width=w, height=h)
    tf = shape.text_frame
    _common.add_paragraph(
        tf,
        card.get("emoji", ""),
        size_pt=theme.EMOJI_SIZE_PT - 8,
        color=theme.TEXT_PRIMARY,
        align=PP_ALIGN.CENTER,
        new=False,
    )
    _common.add_paragraph(
        tf,
        card.get("title", ""),
        size_pt=theme.TEXT_SIZE_PT,
        color=theme.TEXT_PRIMARY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    _common.add_paragraph(
        tf,
        card.get("description", ""),
        size_pt=theme.TEXT_SIZE_PT - 2,
        color=theme.TEXT_SECONDARY,
        align=PP_ALIGN.CENTER,
    )
