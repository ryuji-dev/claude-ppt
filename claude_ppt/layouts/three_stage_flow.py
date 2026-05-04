"""three-stage-flow 레이아웃 — 정확히 3단계 발전 과정.

content 스키마:
    {"stages": [
       {"title": "1단계", "description": "..."},
       {"title": "2단계", "description": "..."},
       {"title": "3단계", "description": "..."}
    ]}
"""

from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt

from .. import theme
from . import _common

STAGE_TOP = Inches(2.6)
STAGE_HEIGHT = Inches(3.6)
STAGE_GAP = Inches(0.6)
STAGE_MARGIN_X = Inches(0.8)

STAGE_COLORS = (theme.SECTION_INTRO, theme.SECTION_1, theme.SECTION_2)


def render(slide: Slide, content: dict, title: str) -> None:
    _common.add_title(slide, title)
    stages = content.get("stages", [])
    if not stages:
        return
    if len(stages) != 3:
        raise ValueError(
            f"three-stage-flow expects exactly 3 stages, got {len(stages)}"
        )

    available = (
        int(_common.SLIDE_WIDTH) - 2 * int(STAGE_MARGIN_X) - int(STAGE_GAP) * 2
    )
    stage_w = available // 3
    centers = []
    for i, stage in enumerate(stages):
        left = int(STAGE_MARGIN_X) + (stage_w + int(STAGE_GAP)) * i
        shape = _common.add_card(
            slide,
            left=left,
            top=int(STAGE_TOP),
            width=stage_w,
            height=int(STAGE_HEIGHT),
            border=STAGE_COLORS[i],
        )
        tf = shape.text_frame
        _common.add_paragraph(
            tf,
            f"0{i + 1}",
            size_pt=theme.LABEL_SIZE_PT + 8,
            color=STAGE_COLORS[i],
            bold=True,
            align=PP_ALIGN.CENTER,
            new=False,
        )
        _common.add_paragraph(
            tf,
            stage.get("title", ""),
            size_pt=theme.TEXT_SIZE_PT + 4,
            color=theme.TEXT_PRIMARY,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        _common.add_paragraph(
            tf,
            stage.get("description", ""),
            size_pt=theme.TEXT_SIZE_PT - 2,
            color=theme.TEXT_SECONDARY,
            align=PP_ALIGN.CENTER,
        )
        centers.append((left, stage_w))

    arrow_y = int(STAGE_TOP) + int(STAGE_HEIGHT) // 2
    for i in range(2):
        start_x = centers[i][0] + centers[i][1]
        end_x = centers[i + 1][0]
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            begin_x=start_x,
            begin_y=arrow_y,
            end_x=end_x,
            end_y=arrow_y,
        )
        connector.line.color.rgb = theme.ACCENT_PURPLE
        connector.line.width = Pt(2.0)
