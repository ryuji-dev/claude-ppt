"""summary-grid 레이아웃 — 6개 요약 카드 + 결론 박스 (마무리 슬라이드 전용).

content 스키마:
    {
      "cards": [{"label": "요약1", "text": "..."}, ...],   # 정확히 6개
      "conclusion": "마무리 결론 텍스트"
    }
"""

from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches

from .. import theme
from . import _common

GRID_TOP = Inches(2.0)
GRID_HEIGHT = Inches(3.4)
GRID_GAP = Inches(0.25)
GRID_MARGIN_X = Inches(0.8)

CONCLUSION_TOP = Inches(5.7)
CONCLUSION_HEIGHT = Inches(1.4)


def render(slide: Slide, content: dict, title: str) -> None:
    _common.add_title(slide, title)
    cards = content.get("cards", [])
    if not cards:
        return

    cols = 3
    rows = (len(cards) + cols - 1) // cols
    available_w = (
        int(_common.SLIDE_WIDTH)
        - 2 * int(GRID_MARGIN_X)
        - int(GRID_GAP) * (cols - 1)
    )
    card_w = available_w // cols
    available_h = int(GRID_HEIGHT) - int(GRID_GAP) * (rows - 1)
    card_h = available_h // rows

    for i, card in enumerate(cards):
        row = i // cols
        col = i % cols
        left = int(GRID_MARGIN_X) + (card_w + int(GRID_GAP)) * col
        top = int(GRID_TOP) + (card_h + int(GRID_GAP)) * row
        shape = _common.add_card(
            slide, left=left, top=top, width=card_w, height=card_h
        )
        tf = shape.text_frame
        _common.add_paragraph(
            tf,
            card.get("label", ""),
            size_pt=theme.LABEL_SIZE_PT,
            color=theme.ACCENT_PURPLE,
            bold=True,
            align=PP_ALIGN.CENTER,
            new=False,
        )
        _common.add_paragraph(
            tf,
            card.get("text", ""),
            size_pt=theme.TEXT_SIZE_PT,
            color=theme.TEXT_PRIMARY,
            align=PP_ALIGN.CENTER,
        )

    conclusion = content.get("conclusion", "")
    if conclusion:
        conc_left = int(GRID_MARGIN_X)
        conc_w = int(_common.SLIDE_WIDTH) - 2 * int(GRID_MARGIN_X)
        shape = _common.add_card(
            slide,
            left=conc_left,
            top=int(CONCLUSION_TOP),
            width=conc_w,
            height=int(CONCLUSION_HEIGHT),
            fill=theme.ACCENT_PURPLE,
            border=theme.ACCENT_PURPLE,
        )
        _common.add_paragraph(
            shape.text_frame,
            conclusion,
            size_pt=theme.TEXT_SIZE_PT + 4,
            color=theme.TEXT_PRIMARY,
            bold=True,
            align=PP_ALIGN.CENTER,
            new=False,
        )
