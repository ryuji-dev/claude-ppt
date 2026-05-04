"""comparison-2col 레이아웃 — 좌·우 2열 비교.

content 스키마:
    {
      "left":  {"title": "Before", "items": ["a", "b"]},
      "right": {"title": "After",  "items": ["c", "d"]}
    }
"""

from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches

from .. import theme
from . import _common

COL_TOP = Inches(2.4)
COL_HEIGHT = Inches(4.4)
COL_GAP = Inches(0.5)
COL_MARGIN_X = Inches(0.8)


def render(slide: Slide, content: dict, title: str) -> None:
    _common.add_title(slide, title)
    left_col = content.get("left", {})
    right_col = content.get("right", {})
    available = int(_common.SLIDE_WIDTH) - 2 * int(COL_MARGIN_X) - int(COL_GAP)
    col_w = available // 2
    _draw_column(slide, int(COL_MARGIN_X), col_w, left_col, theme.ACCENT_PURPLE)
    _draw_column(
        slide,
        int(COL_MARGIN_X) + col_w + int(COL_GAP),
        col_w,
        right_col,
        theme.ACCENT_BLUE,
    )


def _draw_column(slide: Slide, left: int, width: int, col: dict, accent) -> None:
    shape = _common.add_card(
        slide, left=left, top=int(COL_TOP), width=width, height=int(COL_HEIGHT)
    )
    tf = shape.text_frame
    _common.add_paragraph(
        tf,
        col.get("title", ""),
        size_pt=theme.LABEL_SIZE_PT + 6,
        color=accent,
        bold=True,
        align=PP_ALIGN.CENTER,
        new=False,
    )
    for item in col.get("items", []):
        _common.add_paragraph(
            tf,
            f"• {item}",
            size_pt=theme.TEXT_SIZE_PT,
            color=theme.TEXT_PRIMARY,
            align=PP_ALIGN.LEFT,
        )
