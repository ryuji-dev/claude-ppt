"""레이아웃 간 공유 헬퍼 — 슬라이드 크기, 타이틀 렌더, 카드 박스 등."""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.slide import Slide
from pptx.util import Inches, Pt

from .. import theme

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

TITLE_TOP = Inches(0.5)
TITLE_HEIGHT = Inches(1.1)
TITLE_LEFT = Inches(0.5)

CONTENT_TOP = Inches(2.0)
CONTENT_HEIGHT = Inches(5.0)
CONTENT_LEFT = Inches(0.8)
CONTENT_WIDTH = Inches(11.733)  # SLIDE_WIDTH - 2 * 0.8


def add_title(slide: Slide, title: str) -> None:
    """슬라이드 상단에 다크 테마 타이틀을 추가한다."""
    box = slide.shapes.add_textbox(
        left=TITLE_LEFT,
        top=TITLE_TOP,
        width=int(SLIDE_WIDTH) - 2 * int(TITLE_LEFT),
        height=TITLE_HEIGHT,
    )
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.name = theme.FONT_FAMILY
    run.font.size = Pt(theme.TITLE_SIZE_PT)
    run.font.bold = True
    run.font.color.rgb = theme.ACCENT_PURPLE


def add_card(
    slide: Slide,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    fill: RGBColor = theme.CARD_BG,
    border: RGBColor = theme.CARD_BORDER,
):
    """둥근 모서리 카드 박스를 그리고 도형을 반환한다."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left=left,
        top=top,
        width=width,
        height=height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1.0)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)
    tf.margin_bottom = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape


def add_paragraph(
    text_frame,
    text: str,
    *,
    size_pt: int,
    color: RGBColor,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.CENTER,
    new: bool = True,
) -> None:
    """text_frame에 한 줄을 추가한다. 첫 호출이면 new=False로 paragraphs[0]을 사용."""
    p = text_frame.add_paragraph() if new else text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = theme.FONT_FAMILY
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
