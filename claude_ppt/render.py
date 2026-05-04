"""outline.json → .pptx 변환기 진입점."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Mapping

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from . import theme
from .layouts import get_renderer

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def render(outline: Mapping, output_path: str | Path) -> Path:
    """Build a .pptx file from an outline mapping and write it to output_path."""
    output_path = Path(output_path)
    pres = Presentation()
    pres.slide_width = SLIDE_WIDTH
    pres.slide_height = SLIDE_HEIGHT

    blank_layout = pres.slide_layouts[6]  # 6 == fully blank in default template

    for slide_def in outline["slides"]:
        slide = pres.slides.add_slide(blank_layout)
        _set_slide_background(slide, theme.BG)
        renderer = get_renderer(slide_def["layout"])
        renderer(slide, slide_def.get("content", {}), slide_def["title"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    pres.save(str(output_path))
    return output_path


def _set_slide_background(slide, color: RGBColor) -> None:
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="claude-ppt-render",
        description="Render outline.json into an editable .pptx file.",
    )
    parser.add_argument("outline", type=Path, help="Path to outline.json")
    parser.add_argument("output", type=Path, help="Path to write the .pptx")
    args = parser.parse_args(argv)

    outline = json.loads(args.outline.read_text(encoding="utf-8"))
    written = render(outline, args.output)
    print(f"Wrote {written} ({len(outline['slides'])} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
